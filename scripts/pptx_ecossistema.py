# Gera a Proposta Comercial em PowerPoint com o Ecossistema Florim interativo e nativo:
# formas editáveis, links entre slides, transição Transformar (Morph) e gatilhos nas subcategorias.
#
# Uso: npm run pptx  (ou: python pptx_ecossistema.py <pasta com dados.json, ícones e logos> <pasta das páginas> <saida.pptx>)
import glob, json, math, sys, os
from lxml import etree
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

ASSETS, PAGINAS, SAIDA = sys.argv[1], sys.argv[2], sys.argv[3]
SEM_GATILHOS = os.environ.get("SEM_GATILHOS") == "1"
SEM_MORPH = os.environ.get("SEM_MORPH") == "1"
D = json.load(open(os.path.join(ASSETS, "dados.json"), encoding="utf-8"))
CATS = {c["id"]: c for c in D["categorias"]}
SLOTS = D["slots"]

# Espaço de desenho 1600 x 900 unidades -> slide 16:9.
U = 7620  # EMU por unidade
def E(v): return Emu(int(round(v * U)))
def pt(px): return Pt(px * 0.6)

S = 0.84                     # escala da roda
CY = 486
CX_HOME, CX_OPEN = 800, 520
R = dict(emblem=150, bandIn=158, bandOut=196, ringIn=202, ringOut=420, ticks=432, fanIn=444, fanOut=726)
STEP, GAP = 12.5, 1.4
TOPO_PROPOSTA = 4            # o ecossistema entra depois da página 4 (igual a src/components/Deck.tsx)

GOLD = ["F9E898", "E2BE60", "C29727"]
WHITE, GRAPH, INK = "FFFFFF", "1F1F1F", "1F1F1F"
TXT2, TXT3 = "C8C8C8", "8F8F8D"

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}

def polar(r, deg):
    a = math.radians(deg)
    return r * math.sin(a), -r * math.cos(a)

def norm(d):
    d = (d + 180) % 360 - 180
    return 180.0 if d == -180 else d

# ---------------------------------------------------------------- formas customizadas
def add_custom(slide, name, cx, cy, rb, path_pts_list, rot, fill_xml, line_xml):
    """Forma cujo retângulo é o quadrado do círculo de raio rb centrado na roda:
    girar a forma = girar em torno do centro da roda (é o que dá o efeito no Morph)."""
    W = 100000
    def P(x, y):  # x,y relativos ao centro em unidades de desenho -> coordenadas do path
        return int((x / (2 * rb) + 0.5) * W), int((y / (2 * rb) + 0.5) * W)
    paths = ""
    for pts in path_pts_list:
        cmds = []
        for i, (x, y) in enumerate(pts):
            px, py = P(x, y)
            cmds.append(f'<a:{"moveTo" if i == 0 else "lnTo"}><a:pt x="{px}" y="{py}"/></a:{"moveTo" if i == 0 else "lnTo"}>')
        paths += f'<a:path w="{W}" h="{W}">' + "".join(cmds) + "<a:close/></a:path>"
    sid = slide.shapes._next_shape_id
    xml = f'''<p:sp xmlns:p="{NS['p']}" xmlns:a="{NS['a']}" xmlns:r="{NS['r']}">
  <p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr><a:xfrm rot="{int(round((rot % 360) * 60000))}"><a:off x="{int((cx - rb) * U)}" y="{int((cy - rb) * U)}"/><a:ext cx="{int(2 * rb * U)}" cy="{int(2 * rb * U)}"/></a:xfrm>
  <a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/><a:pathLst>{paths}</a:pathLst></a:custGeom>
  {fill_xml}{line_xml}</p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="pt-BR"/></a:p></p:txBody></p:sp>'''
    el = parse_xml(xml)
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]

def sector_pts(ri, ro, a, b, step=1.0):
    n = max(2, int(abs(b - a) / step) + 1)
    outer = [polar(ro, a + (b - a) * i / (n - 1)) for i in range(n)]
    inner = [polar(ri, b - (b - a) * i / (n - 1)) for i in range(n)]
    return outer + inner

def solid(c, alpha=None):
    a = f'<a:alpha val="{int(alpha * 1000)}"/>' if alpha is not None else ""
    return f'<a:solidFill><a:srgbClr val="{c}">{a}</a:srgbClr></a:solidFill>'

def line(c, w=0.9, alpha=None):
    return f'<a:ln w="{int(w * 0.6 * 12700)}">{solid(c, alpha)}</a:ln>'

NOLINE = "<a:ln><a:noFill/></a:ln>"
GOLD_LIN = ('<a:gradFill rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="F9E898"/></a:gs>'
            '<a:gs pos="50000"><a:srgbClr val="E2BE60"/></a:gs><a:gs pos="100000"><a:srgbClr val="C29727"/></a:gs></a:gsLst>'
            '<a:lin ang="5400000" scaled="0"/></a:gradFill>')
SECTOR_FILL = ('<a:gradFill rotWithShape="0"><a:gsLst><a:gs pos="0"><a:srgbClr val="121212"/></a:gs>'
               '<a:gs pos="50000"><a:srgbClr val="161616"/></a:gs><a:gs pos="100000"><a:srgbClr val="2A2A28"/></a:gs></a:gsLst>'
               '<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path></a:gradFill>')
CARD_FILL = ('<a:gradFill rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="2E2E2C"/></a:gs>'
             '<a:gs pos="45000"><a:srgbClr val="242422"/></a:gs><a:gs pos="100000"><a:srgbClr val="161616"/></a:gs></a:gsLst>'
             '<a:lin ang="3600000" scaled="0"/></a:gradFill>')

def set_fill_xml(shape, fill_xml, line_xml=NOLINE):
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:ln"):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    for x in (fill_xml, line_xml):
        el = parse_xml(f'<root xmlns:a="{NS["a"]}">{x}</root>')[0]
        spPr.append(el)

# ---------------------------------------------------------------- texto
def text(slide, name, x, y, w, h, runs, align="l", anchor="t", rot=0, wrap=True):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    first = True
    for para in runs:  # lista de parágrafos; cada parágrafo é lista de (texto, dict)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        for t, st in para:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = st.get("font", "Poppins")
            f.size = pt(st.get("size", 16))
            f.bold = st.get("bold", False)
            f.italic = st.get("italic", False)
            f.color.rgb = RGBColor.from_string(st.get("color", WHITE))
            if st.get("spacing"):
                r._r.get_or_add_rPr().set("spc", str(int(st["spacing"] * 100)))
            if st.get("caps"):
                r._r.get_or_add_rPr().set("cap", "all")
        if para and para[0][1].get("line"):
            p.line_spacing = para[0][1]["line"]
        if para and para[0][1].get("after"):
            p.space_after = pt(para[0][1]["after"])
    if rot:
        tb.rotation = rot
    # caixa de texto clicável na área inteira (preenchimento invisível)
    set_fill_xml(tb, solid("000000", 0))
    return tb

def pill(slide, name, x, y, w, h, label, size=10.5, fill=GOLD_LIN, color=INK, outline=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    s.name = name
    s.adjustments[0] = 0.5
    set_fill_xml(s, fill, outline or NOLINE)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = E(10); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.name = "Poppins SemiBold"; r.font.size = pt(size); r.font.color.rgb = RGBColor.from_string(color)
    r._r.get_or_add_rPr().set("spc", "160"); r._r.get_or_add_rPr().set("cap", "all")
    return s

def card(slide, name, x, y, w, h, radius=0.12):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    s.name = name
    s.adjustments[0] = radius
    set_fill_xml(s, CARD_FILL, line("FFFFFF", 0.8, 8))
    s.text_frame.text = ""
    return s

def link_slide(shape, target):
    shape.click_action.target_slide = target

def link_url(shape, url):
    shape.click_action.hyperlink.address = url

# ---------------------------------------------------------------- fundo
def make_background(path):
    W, H = 1920, 1080
    img = Image.new("RGB", (W, H), (5, 5, 5))
    over = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(over)
    for x in range(-40, W + 40, 86):
        edge = abs((x / W) - 0.5) * 2  # 0 no centro, 1 nas bordas
        a = int(18 + 60 * edge ** 1.6)
        d.rectangle([x, 0, x + 34, H], fill=(194, 151, 39, a))
    over = over.filter(ImageFilter.GaussianBlur(14))
    glow = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    ImageDraw.Draw(glow).ellipse([W * 0.2, -H * 0.55, W * 0.8, H * 0.35], fill=(194, 151, 39, 40))
    glow = glow.filter(ImageFilter.GaussianBlur(120))
    vign = Image.new("L", (W, H), 0)
    ImageDraw.Draw(vign).ellipse([W * 0.12, H * 0.05, W * 0.88, H * 1.05], fill=255)
    vign = vign.filter(ImageFilter.GaussianBlur(160))
    img = Image.alpha_composite(img.convert("RGBA"), over)
    img = Image.alpha_composite(img, glow)
    dark = Image.new("RGBA", (W, H), (5, 5, 5, 200))
    img = Image.composite(dark, img, vign)
    img.convert("RGB").save(path, quality=92)

# ---------------------------------------------------------------- slide do ecossistema
def build_eco_slide(slide, sel, slides_by_key, continuar):
    slot = next((s for s in SLOTS if s["id"] == sel), None)
    rot = norm(90 - slot["mid"]) if slot else 0.0
    cx = CX_OPEN if slot else CX_HOME
    cat = CATS[sel] if sel else None
    home = slides_by_key["home"]

    bg = slide.shapes.add_picture(os.path.join(ASSETS, "fundo.jpg"), 0, 0, E(1600), E(900))
    bg.name = "fundo"

    # Pétalas recolhidas (escondidas sob a roda) e abertas na categoria escolhida.
    petal_shapes = []
    for s in SLOTS:
        c = CATS[s["id"]]
        n = len(c["subs"])
        hw = (STEP - GAP) / 2
        pts = sector_pts(R["fanIn"] * S, R["fanOut"] * S, -hw, hw, 0.5)
        for i in range(n):
            nm = f"!!petala-{c['id']}-{i}"
            if sel == c["id"]:
                off = -n * STEP / 2 + STEP * (i + 0.5)
                shp = add_custom(slide, nm, cx, CY, R["fanOut"] * S, [pts], 90 + off,
                                 solid(GRAPH), line("E2BE60", 1, 55))
                petal_shapes.append((i, shp, off))
            else:
                k = 0.5
                scaled = [(x * k, y * k) for x, y in pts]
                add_custom(slide, nm, cx, CY, R["fanOut"] * S * k, [scaled], s["mid"] + rot,
                           solid(GRAPH), line("E2BE60", 1, 55))

    # pescoço ligando a categoria ao leque
    if slot:
        n = len(cat["subs"])
        half_sel = (slot["b"] - slot["a"]) / 2 - 0.6
        half_fan = n * STEP / 2
        r1, r2 = R["ringOut"] * S + 2, R["fanIn"] * S - 1
        neck = [polar(r1, 90 - half_sel)] + [polar(r2, 90 - half_fan + 2 * half_fan * t / 20) for t in range(21)] + \
               [polar(r1, 90 + half_sel - 2 * half_sel * t / 10) for t in range(11)]
        add_custom(slide, "pescoco", cx, CY, r2, [neck], 0, solid("E2BE60", 9), line("E2BE60", 0.8, 35))

    # anel de marcações
    tk = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(cx - R["ticks"] * S), E(CY - R["ticks"] * S), E(2 * R["ticks"] * S), E(2 * R["ticks"] * S))
    tk.name = "!!marcacoes"
    set_fill_xml(tk, "<a:noFill/>", f'<a:ln w="{int(9 * 12700)}">{solid("E2BE60", 22)}<a:prstDash val="sysDot"/></a:ln>')
    tk.text_frame.text = ""

    # faixas das frentes
    for bnd in D["bands"]:
        add_custom(slide, "!!faixa-" + bnd["nome"], cx, CY, R["bandOut"] * S,
                   [sector_pts(R["bandIn"] * S, R["bandOut"] * S, bnd["a"] + 0.4, bnd["b"] - 0.4)], rot,
                   solid("0D0D0D"), line("E2BE60", 0.8, 30))

    # setores
    for s in SLOTS:
        c = CATS[s["id"]]
        is_sel = sel == c["id"]
        shp = add_custom(slide, "!!setor-" + c["id"], cx, CY, R["ringOut"] * S,
                         [sector_pts(R["ringIn"] * S, R["ringOut"] * S, s["a"] + 0.45, s["b"] - 0.45)], rot,
                         GOLD_LIN if is_sel else SECTOR_FILL, line("E2BE60", 0.9, 26))
        dest = home if is_sel else slides_by_key[c["id"]]
        link_slide(shp, dest)
        if sel and not is_sel:
            set_fill_xml(shp, SECTOR_FILL.replace('val="2A2A28"', 'val="1C1C1B"'), line("E2BE60", 0.9, 14))

        big = c["frente"] != "servicos"
        lx, ly = polar((311 if big else 318) * S, s["mid"] + rot)
        lx += cx; ly += CY
        icon = 40 if big else 30
        lines = c.get("linhas") or [c["rotulo"]]
        block_h = icon * S + 8 + len(lines) * 19 + (20 if c.get("chamada") else 0)
        top = ly - block_h / 2
        ic = slide.shapes.add_picture(os.path.join(ASSETS, f"ico-{c['id']}-{'escuro' if is_sel else 'ouro'}.png"),
                                      E(lx - icon * S / 2), E(top), E(icon * S), E(icon * S))
        ic.name = "!!icone-" + c["id"]
        link_slide(ic, dest)
        dim = sel and not is_sel
        col = INK if is_sel else ("8A8A88" if dim else WHITE)
        paras = [[(l, {"size": 18 if big else 15, "font": "Poppins SemiBold" if big else "Poppins Medium", "color": col, "line": 0.95})] for l in lines]
        if c.get("chamada"):
            paras.append([(c["chamada"], {"size": 12, "font": "Poppins Light", "color": "4A3C14" if is_sel else TXT3})])
        tb = text(slide, "!!rotulo-" + c["id"], lx - 80, top + icon * S + 6, 160, block_h - icon * S, paras, align="c")
        link_slide(tb, dest)

    # marcação dos projetos complementares (só na roda fechada)
    if not slot and D["comp"]:
        a, b = D["comp"]["a"], D["comp"]["b"]
        arc = [polar(R["ringOut"] * S + 10, a + 1 + (b - a - 2) * t / 40) for t in range(41)]
        arc2 = [polar(R["ringOut"] * S + 12, b - 1 - (b - a - 2) * t / 40) for t in range(41)]
        add_custom(slide, "complementares", cx, CY, R["ringOut"] * S + 12, [arc + arc2], 0, solid("E2BE60"), NOLINE)
        mx, my = polar(R["ringOut"] * S + 34, (a + b) / 2)
        text(slide, "complementares-txt", cx + mx - 90, CY + my - 10, 180, 20,
             [[("PROJETOS COMPLEMENTARES", {"size": 10.5, "font": "Poppins SemiBold", "color": "E2BE60", "spacing": 2.4})]],
             align="c", rot=(a + b) / 2)

    # brasão
    er = R["emblem"] * S
    disc = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(cx - er), E(CY - er), E(2 * er), E(2 * er))
    disc.name = "!!brasao-disco"
    set_fill_xml(disc, '<a:gradFill rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="2E2E2C"/></a:gs>'
                       '<a:gs pos="100000"><a:srgbClr val="0A0A0A"/></a:gs></a:gsLst><a:path path="circle">'
                       '<a:fillToRect l="50000" t="40000" r="50000" b="60000"/></a:path></a:gradFill>',
                 line("E2BE60", 1.2, 60))
    disc.text_frame.text = ""
    em = slide.shapes.add_picture(os.path.join(ASSETS, "florim-emblema.png"), E(cx - 112 * S), E(CY - 109 * S), E(224 * S), E(218 * S))
    em.name = "!!brasao"
    for sh in (disc, em):
        link_slide(sh, home)

    # pétalas: rótulos (por cima) e realces
    desc_groups, highlight, petal_hits = [], [], []
    if slot:
        n = len(cat["subs"])
        hw = (STEP - GAP) / 2
        hl_pts = sector_pts(R["fanIn"] * S, R["fanOut"] * S, -hw, hw, 0.5)
        for i, shp, off in petal_shapes:
            h = add_custom(slide, f"realce-{i}", cx, CY, R["fanOut"] * S, [hl_pts], 90 + off,
                           GOLD_LIN, NOLINE)
            highlight.append(h)
        for i, shp, off in petal_shapes:
            mr = (R["fanIn"] + R["fanOut"]) / 2 * S
            px, py = polar(mr, 90 + off)
            w = (R["fanOut"] - R["fanIn"]) * S - 26
            sub = cat["subs"][i]
            tb = text(slide, f"rotulo-petala-{i}", cx + px - w / 2, CY + py - 20, w, 40,
                      [[(f"{i + 1:02d}   ", {"size": 13, "font": "Poppins SemiBold", "color": "E2BE60"}),
                        (sub["titulo"], {"size": 15.5, "font": "Poppins Medium", "color": WHITE})]],
                      anchor="m", rot=off)
            petal_hits.append((i, shp, tb))

    # ---------------- lateral esquerda
    if not slot:
        text(slide, "titulo", 70, 300, 360, 200,
             [[("Ecossistema", {"size": 60, "font": "Poppins Light", "color": "E2BE60", "line": 0.95})],
              [("Florim", {"size": 64, "font": "Poppins SemiBold", "italic": True, "color": "E2BE60", "line": 0.95})]])
        text(slide, "subtitulo", 70, 470, 330, 30, [[("Um só parceiro, soluções completas.", {"size": 17, "color": WHITE})]])
        text(slide, "texto", 70, 506, 300, 110,
             [[("Engenharia, representações, cursos e consultoria sob a mesma marca. Mais de ", {"size": 14, "font": "Poppins Light", "color": TXT2, "line": 1.3}),
               ("10 anos", {"size": 14, "font": "Poppins SemiBold", "color": WHITE}),
               (" entregando projetos 360°.", {"size": 14, "font": "Poppins Light", "color": TXT2})]])
        ar = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(70), E(640), E(40), E(40)); ar.name = "seta"
        set_fill_xml(ar, GOLD_LIN); ar.text_frame.text = "›"
        r = ar.text_frame.paragraphs[0].runs[0]; r.font.size = pt(26); r.font.color.rgb = RGBColor.from_string(INK); r.font.name = "Poppins"
        ar.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER; ar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        ar.text_frame.margin_top = ar.text_frame.margin_bottom = 0
        text(slide, "dica", 124, 650, 260, 24, [[("Clique em uma área para explorar", {"size": 13, "color": TXT2})]])
    else:
        b = pill(slide, "voltar", 64, 92, 190, 40, "‹  Todas as áreas", size=11, fill=solid("1F1F1F"), color=WHITE,
                 outline=line("E2BE60", 1, 40))
        link_slide(b, home)

    # ---------------- lateral direita
    if not slot:
        x0 = 1230
        pill(slide, "badge-frentes", x0, 248, 150, 26, "Como atuamos", size=10)
        text(slide, "titulo-frentes", x0, 284, 340, 60,
             [[("Quatro ", {"size": 44, "font": "Poppins Light", "color": WHITE}),
               ("frentes", {"size": 44, "font": "Poppins SemiBold", "italic": True, "color": WHITE})]])
        itens = [("servicos", "Serviços de engenharia", "9 disciplinas", "agrimensura"),
                 ("representacoes", "Representações", "4 soluções", "representacoes"),
                 ("consultoria", "Consultoria", "Inteligência de mercado", "consultoria"),
                 ("cursos", "Cursos NR", "Aulas gratuitas", "cursos")]
        for k, (f, nome, sub, first) in enumerate(itens):
            y = 360 + k * 76
            cd = card(slide, f"frente-{f}", x0, y, 320, 66, radius=0.28)
            link_slide(cd, slides_by_key[first])
            tb = text(slide, f"frente-txt-{f}", x0 + 18, y + 12, 290, 44,
                      [[(f"{k + 1:02d}   ", {"size": 12, "font": "Poppins SemiBold", "color": "E2BE60"}),
                        (nome, {"size": 14, "font": "Poppins SemiBold", "color": "E2BE60", "caps": True, "spacing": 0.6})],
                       [("        " + sub, {"size": 12.5, "font": "Poppins Light", "color": TXT2})]])
            link_slide(tb, slides_by_key[first])
    else:
        x0, y0, w, h = 1172, 196, 384, 560
        kicker = D["frentes"][cat["frente"]]["nome"] + (" · " + cat["grupo"] if cat.get("grupo") else "")

        def panel(prefix, title_runs, body, cta_url, cta_label, number=None):
            shapes = [card(slide, prefix + "cartao", x0, y0, w, h, radius=0.08)]
            shapes.append(pill(slide, prefix + "rotulo", x0 + 30, y0 + 32, min(320, 40 + len(kicker) * 7.4), 26, kicker, size=9.5))
            y = y0 + 84
            if number is not None:
                shapes.append(text(slide, prefix + "numero", x0 + 30, y, 200, 64,
                                   [[(f"{number:02d}", {"size": 54, "font": "Poppins SemiBold", "italic": True, "color": "E2BE60"})]]))
                y += 68
            shapes.append(text(slide, prefix + "titulo", x0 + 30, y, w - 60, 110, [title_runs]))
            longo = sum(len(t) for t, _ in title_runs) > 22
            y += 62 if not longo else 100
            shapes.append(text(slide, prefix + "texto", x0 + 30, y, w - 60, 150,
                               [[(body, {"size": 15, "font": "Poppins Light", "color": TXT2, "line": 1.35})]]))
            cta = pill(slide, prefix + "cta", x0 + 30, y0 + h - 116, 240, 46, cta_label + "  →", size=11)
            link_url(cta, cta_url)
            shapes.append(cta)
            if cat["acao"]["tipo"] == "contato":
                m = text(slide, prefix + "email", x0 + 30, y0 + h - 56, w - 60, 22,
                         [[("ou escreva para ", {"size": 12, "color": TXT3}),
                           (D["config"]["email"], {"size": 12, "color": TXT2})]])
                link_url(m, D["email"][cat["id"]])
                shapes.append(m)
            return shapes

        url = D["config"]["cursosUrl"] if cat["acao"]["tipo"] == "cursos" else D["contato"][cat["id"]][0]
        panel("resumo-", [(cat["rotulo"], {"size": 40, "font": "Poppins SemiBold", "italic": True, "color": "E2BE60", "line": 1.0})],
              cat["resumo"] + f"\n\nClique nas {len(cat['subs'])} opções do leque para ver os detalhes.", url, cat["acao"]["texto"])
        for i, sub in enumerate(cat["subs"]):
            u = D["config"]["cursosUrl"] if cat["acao"]["tipo"] == "cursos" else D["contato"][cat["id"]][i + 1]
            size = 30 if any(len(wd) > 13 for wd in sub["titulo"].split()) else 36
            shapes = panel(f"desc{i}-", [(sub["titulo"], {"size": size, "font": "Poppins Medium", "color": WHITE, "line": 1.0})],
                           sub["texto"], u, cat["acao"]["texto"], number=i + 1)
            grp = slide.shapes.add_group_shape(shapes)
            grp.name = f"descricao-{i}"
            desc_groups.append(grp)

    # ---------------- topo
    pill(slide, "badge-topo", 48, 20, 132, 26, "Ecossistema", size=10)
    e2 = slide.shapes.add_picture(os.path.join(ASSETS, "florim-emblema.png"), E(712), E(12), E(42), E(41)); e2.name = "logo-brasao"
    wm = slide.shapes.add_picture(os.path.join(ASSETS, "florim-wordmark.png"), E(758), E(20), E(140), E(25.5)); wm.name = "logo-nome"
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(48), E(64), E(1504), E(1)); ln.name = "linha-topo"
    set_fill_xml(ln, solid("FFFFFF", 16)); ln.text_frame.text = ""
    c = pill(slide, "continuar", 1352, 18, 200, 30, "Continuar proposta  ›", size=10)
    link_slide(c, continuar)
    gl = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, E(895), E(1600), E(5)); gl.name = "linha-ouro"
    set_fill_xml(gl, '<a:gradFill><a:gsLst><a:gs pos="0"><a:srgbClr val="C29727"/></a:gs><a:gs pos="50000"><a:srgbClr val="F9E898"/></a:gs>'
                     '<a:gs pos="100000"><a:srgbClr val="C29727"/></a:gs></a:gsLst><a:lin ang="0" scaled="0"/></a:gradFill>')
    gl.text_frame.text = ""

    return desc_groups, highlight, petal_hits

# ---------------------------------------------------------------- animações (gatilhos)
class Ids:
    def __init__(self): self.n = 1
    def __call__(self):
        self.n += 1
        return self.n

def eff(ids, spid, entr, dur=320):
    if entr:
        return (f'<p:par><p:cTn id="{ids()}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="withEffect">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{ids()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{ids()}" dur="{dur}"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')
    return (f'<p:par><p:cTn id="{ids()}" presetID="10" presetClass="exit" presetSubtype="0" fill="hold" nodeType="withEffect">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn id="{ids()}" dur="150"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            f'<p:set><p:cBhvr><p:cTn id="{ids()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="149"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:to><p:strVal val="hidden"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')

def seq(ids, trigger_spid, effects):
    if effects:
        effects = [effects[0].replace('nodeType="withEffect"', 'nodeType="clickEffect"', 1)] + effects[1:]
    return (f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{ids()}" restart="whenNotActive" fill="hold" evtFilter="cancelBubble" nodeType="interactiveSeq">'
            f'<p:stCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{trigger_spid}"/></p:tgtEl></p:cond></p:stCondLst>'
            f'<p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync><p:childTnLst>'
            f'<p:par><p:cTn id="{ids()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{ids()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            + "".join(effects) +
            f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
            f'<p:nextCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{trigger_spid}"/></p:tgtEl></p:cond></p:nextCondLst></p:seq>')

def add_triggers(slide, desc_groups, highlight, petal_hits, primer):
    ids = Ids()
    seqs = []
    allspid = [g.shape_id for g in desc_groups] + [h.shape_id for h in highlight]
    # Primeira sequência: só entradas, num alvo que nunca é clicado. Faz o PowerPoint começar com tudo oculto.
    seqs.append(seq(ids, primer.shape_id, [eff(ids, s, True, 1) for s in allspid]))
    for i, shp, tb in petal_hits:
        effects = []
        for j in range(len(desc_groups)):
            if j != i:
                effects.append(eff(ids, desc_groups[j].shape_id, False))
                effects.append(eff(ids, highlight[j].shape_id, False))
        effects.append(eff(ids, highlight[i].shape_id, True, 250))
        effects.append(eff(ids, desc_groups[i].shape_id, True, 320))
        for target in (shp, tb):
            seqs.append(seq(ids, target.shape_id, effects))
            effects = [e for e in effects]  # mesmos efeitos para o rótulo
    # ids precisam ser únicos: regera cada sequência com novos ids
    xml = (f'<p:timing xmlns:p="{NS["p"]}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
           f'<p:childTnLst>{"".join(seqs)}</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    el = etree.fromstring(xml)
    # renumera ids de cTn para garantir unicidade
    for k, ctn in enumerate(el.iter(qn("p:cTn")), start=1):
        ctn.set("id", str(k))
    slide._element.append(el)

def add_transition(slide, morph=True, adv_click=False):
    adv = "" if adv_click else ' advClick="0"'
    if morph:
        xml = (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:p="{NS["p"]}">'
               f'<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
               f'<p:transition spd="slow"{adv} xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" p14:dur="1100">'
               f'<p159:morph option="byObject"/></p:transition></mc:Choice>'
               f'<mc:Fallback><p:transition spd="slow"{adv}><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    else:
        xml = f'<p:transition xmlns:p="{NS["p"]}" spd="med"{adv}><p:fade/></p:transition>'
    el = etree.fromstring(xml)
    sld = slide._element
    clr = sld.find(qn("p:clrMapOvr"))
    if clr is not None:
        clr.addnext(el)
    else:
        sld.find(qn("p:cSld")).addnext(el)
    return el

# ---------------------------------------------------------------- montagem
def main():
    make_background(os.path.join(ASSETS, "fundo.jpg"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(1600), E(900)
    blank = prs.slide_layouts[6]

    imgs = sorted(glob.glob(os.path.join(PAGINAS, "pagina-*.jpg")))

    def page_slide(img):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(img, 0, 0, E(1600), E(900)).name = "pagina"
        add_transition(s, morph=False)
        return s

    before = [page_slide(imgs[i]) for i in range(TOPO_PROPOSTA)]
    keys = ["home"] + [s["id"] for s in SLOTS]
    eco = {k: prs.slides.add_slide(blank) for k in keys}
    after = [page_slide(imgs[i]) for i in range(TOPO_PROPOSTA, len(imgs))]
    continuar = after[0]

    for k, slide in eco.items():
        sel = None if k == "home" else k
        desc, hl, hits = build_eco_slide(slide, sel, eco, continuar)
        add_transition(slide, morph=not SEM_MORPH, adv_click=False)
        if sel:
            slide._element.set("show", "0")  # oculto: só se chega por clique
            primer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(0), E(0), E(2), E(2))
            primer.name = "gatilho-inicial"
            if not SEM_GATILHOS:
                add_triggers(slide, desc, hl, hits, primer)

    prs.core_properties.title = "Proposta Comercial Florim Engenharia"
    prs.core_properties.author = "Florim Engenharia"
    prs.save(SAIDA)
    print("salvo", SAIDA, len(prs.slides), "slides")

main()
